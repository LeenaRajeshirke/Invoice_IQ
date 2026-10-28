import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas.invoice import CURRENCY_CODES, InvoiceFields


class TabularFileError(ValueError):
    """A CSV or Excel file cannot be interpreted as invoice records."""


class TabularInvoiceRecord:
    def __init__(self, fields: InvoiceFields, warnings: list[str]) -> None:
        self.fields = fields
        self.warnings = warnings


HEADER_ALIASES = {
    "invoice_number": ("invoice_number", "invoice_no", "invoice_num", "invoice_id", "invoice"),
    "vendor_name": ("vendor_name", "vendor", "supplier_name", "supplier", "seller"),
    "vendor_address": ("vendor_address", "supplier_address", "seller_address"),
    "vendor_tax_id": ("vendor_tax_id", "supplier_tax_id", "tax_id", "gstin"),
    "customer_name": ("customer_name", "customer", "buyer_name", "buyer"),
    "invoice_date": ("invoice_date", "date_of_invoice", "issued_date", "date_issued"),
    "due_date": ("due_date", "payment_due_date"),
    "currency": ("currency", "currency_code"),
    "subtotal": ("subtotal", "sub_total", "net_amount"),
    "tax_amount": ("tax_amount", "tax", "vat", "gst"),
    "discount_amount": ("discount_amount", "discount"),
    "total_amount": ("total_amount", "total", "grand_total", "amount_due"),
    "amount_paid": ("amount_paid", "paid"),
    "balance_due": ("balance_due", "outstanding_amount"),
}
HEADER_TO_FIELD = {
    alias: field for field, aliases in HEADER_ALIASES.items() for alias in aliases
}
DATE_FORMATS = (
    "%Y/%m/%d",
    "%B %d, %Y",
    "%b %d, %Y",
    "%d %B %Y",
    "%d %b %Y",
)


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            pass

    match = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text)
    if not match:
        return None
    first, second, year = (int(part) for part in match.groups())
    if first <= 12 and second <= 12 and first != second:
        return None
    month, day = (second, first) if first > 12 else (first, second)
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _parse_amount(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        amount = value
    elif isinstance(value, (int, float)):
        amount = Decimal(str(value))
    else:
        text = str(value).strip()
        if not text:
            return None
        text = re.sub(r"^(?:[$\u20ac\u00a3]\s*)", "", text)
        text = re.sub(r"\s*(?:AUD|CAD|CHF|CNY|EUR|GBP|INR|JPY|NZD|USD)$", "", text, flags=re.I)
        if text.startswith("(") and text.endswith(")"):
            text = f"-{text[1:-1]}"
        if re.fullmatch(r"[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
            text = text.replace(",", "")
        elif not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
            return None
        try:
            amount = Decimal(text)
        except InvalidOperation:
            return None
    if (
        not amount.is_finite()
        or amount.as_tuple().exponent < -2
        or abs(amount) > Decimal("9999999999.99")
    ):
        return None
    return amount


def _record_from_row(row: dict[str, Any]) -> TabularInvoiceRecord:
    values = {_header_key(key): value for key, value in row.items() if key is not None}
    fields: dict[str, Any] = {}
    warnings: list[str] = []
    for column, value in values.items():
        field = HEADER_TO_FIELD.get(column)
        if field is None or value is None or str(value).strip() == "":
            continue
        if field in {"invoice_date", "due_date"}:
            parsed = _parse_date(value)
            if parsed is None:
                warnings.append(f"{field.replace('_', ' ').title()} is invalid or ambiguous; review is required.")
            else:
                fields[field] = parsed
        elif field in {
            "subtotal",
            "tax_amount",
            "discount_amount",
            "total_amount",
            "amount_paid",
            "balance_due",
        }:
            amount = _parse_amount(value)
            if amount is None:
                warnings.append(f"{field.replace('_', ' ').title()} is not a finite numeric amount; review is required.")
            else:
                fields[field] = amount
        elif field == "currency":
            currency = str(value).strip().upper()
            if currency not in CURRENCY_CODES:
                warnings.append("Currency code is unsupported or unclear; review is required.")
            else:
                fields[field] = currency
        else:
            fields[field] = str(value).strip()
    return TabularInvoiceRecord(fields=InvoiceFields(**fields), warnings=warnings)


def _extract_rows(rows: list[tuple[Any, ...]]) -> list[TabularInvoiceRecord]:
    if not rows:
        raise TabularFileError("The file contains no invoice records.")
    headers = [_header_key(value) for value in rows[0]]
    if not any(header in HEADER_TO_FIELD for header in headers):
        raise TabularFileError(
            "No recognized invoice column headers were found. Include fields such as invoice_number, vendor_name, and total_amount."
        )
    records: list[TabularInvoiceRecord] = []
    for values in rows[1:]:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        records.append(_record_from_row(dict(zip(headers, values))))
    if not records:
        raise TabularFileError("The file contains headers but no invoice records.")
    return records


def extract_csv_invoices(file_path: Path) -> list[TabularInvoiceRecord]:
    try:
        text = file_path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeDecodeError) as error:
        raise TabularFileError("The CSV is corrupt or is not valid UTF-8 text.") from error
    if not text.strip():
        raise TabularFileError("The CSV file is empty.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as error:
        raise TabularFileError("The CSV contents could not be parsed.") from error
    return _extract_rows(rows)


def extract_excel_invoices(file_path: Path) -> list[TabularInvoiceRecord]:
    try:
        with zipfile.ZipFile(file_path) as archive:
            if archive.testzip() is not None:
                raise TabularFileError("The Excel workbook is corrupt.")
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as error:
        raise TabularFileError("The Excel workbook is corrupt or unreadable.") from error
    try:
        records: list[TabularInvoiceRecord] = []
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            try:
                records.extend(_extract_rows(rows))
            except TabularFileError:
                continue
        if not records:
            raise TabularFileError("The Excel workbook contains no invoice records.")
        return records
    finally:
        workbook.close()
